#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Parse WorkBuddy credit-consumption exports (积分消耗明细) into per-task usage.

Source: the export button on WorkBuddy's usage page (浏览器 -> 积分消耗明细 -> 导出).
Export files are auto-discovered in the user's Downloads folder (or the bridge
dir) by name pattern, so the user only needs to click "导出" now and then:

  *积分消耗明细*.csv / .xlsx   (also matches generic usage/消耗 names)

Aggregation (group by the 请求/title column = the conversation/task name):
  credits_used = sum(cost)      last_activity = max(time)
  status: "running" if last activity within RUNNING_HOURS else "done"
    - done  -> progress 1.0 (not running any more = finished)
    - running -> progress stays 0 and the cover shows credits used instead

Env:  WB_RUNNING_HOURS (default 24)  - how recent activity counts as "running"
"""
import csv
import glob
import os
from datetime import datetime

HOME = os.path.expanduser("~")
DOWNLOADS = os.path.join(HOME, "Downloads")
BRIDGE_DIR = os.path.dirname(os.path.abspath(__file__))
RUNNING_HOURS = float(os.environ.get("WB_RUNNING_HOURS", "24"))
MAX_RUNNING = 5          # cover shows at most 5 running rows
MAX_DONE = 6             # cover shows at most 6 done rows at the bottom

_NAME_PATTERNS = ["*积分消耗明细*", "*消耗*", "*usage*", "*credits*"]
_EXTS = (".csv", ".xlsx")


def find_exports():
    """All plausible export files, newest first."""
    out = []
    for d in (DOWNLOADS, BRIDGE_DIR):
        for pat in _NAME_PATTERNS:
            for ext in _EXTS:
                out.extend(glob.glob(os.path.join(d, pat + ext)))
    seen, files = set(), []
    for f in sorted(set(out), key=os.path.getmtime, reverse=True):
        if f not in seen:
            seen.add(f)
            files.append(f)
    return files


def _rows_csv(path):
    rows = []
    for enc in ("utf-8-sig", "gbk", "utf-8"):
        try:
            with open(path, "r", encoding=enc, newline="") as f:
                rows = list(csv.reader(f))
            if rows:
                break
        except Exception:
            continue
    return rows


def _rows_xlsx(path):
    try:
        from openpyxl import load_workbook
    except Exception:
        return []
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        return [[("" if c is None else str(c)) for c in r]
                for r in ws.iter_rows(values_only=True)]
    except Exception:
        return []


def _parse_time(s):
    s = str(s).strip()
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y/%m/%d %H:%M:%S",
                "%Y-%m-%d %H:%M", "%Y/%m/%d %H:%M", "%Y-%m-%d", "%Y/%m/%d"):
        try:
            return datetime.strptime(s[:19], fmt)
        except Exception:
            pass
    return None


def _parse_cost(s):
    s = str(s).replace(",", "").strip()
    try:
        return float(s)
    except Exception:
        return 0.0


def _detect_cols(header):
    h = [str(x).strip().lower() for x in header]

    def find(*keys):
        for i, x in enumerate(h):
            if any(k in x for k in keys):
                return i
        return None

    return {
        "time": find("时间", "time"),
        "title": find("请求", "标题", "title", "request", "任务"),
        "cost": find("积分", "消耗", "cost", "credit"),
    }


def aggregate(paths=None):
    """Parse all export files -> list of task dicts (newest activity first)."""
    files = paths or find_exports()
    if not files:
        return []
    recs = {}
    newest_file = None
    for fp in files:
        rows = _rows_csv(fp) if fp.lower().endswith(".csv") else _rows_xlsx(fp)
        if len(rows) < 2:
            continue
        cols = _detect_cols(rows[0])
        if cols["title"] is None or cols["cost"] is None:
            continue
        if newest_file is None or os.path.getmtime(fp) > os.path.getmtime(newest_file):
            newest_file = fp
        for r in rows[1:]:
            try:
                ti = cols["title"]
                if ti >= len(r):
                    continue
                title = str(r[ti]).strip()
                if not title:
                    continue
                cost = _parse_cost(r[cols["cost"]]) if cols["cost"] < len(r) else 0.0
                ts = None
                if cols["time"] is not None and cols["time"] < len(r):
                    ts = _parse_time(r[cols["time"]])
            except Exception:
                continue
            g = recs.setdefault(title, {"credits": 0.0, "last": None})
            g["credits"] += cost
            if ts and (g["last"] is None or ts > g["last"]):
                g["last"] = ts
    now = datetime.now()
    tasks = []
    for name, g in recs.items():
        lt = g["last"]
        hours = (now - lt).total_seconds() / 3600.0 if lt else 1e9
        running = hours <= RUNNING_HOURS
        tasks.append({
            "name": name,
            "status": "running" if running else "done",
            "progress": 0.0 if running else 1.0,
            "credits": round(g["credits"], 2),
            "date": lt.strftime("%Y-%m-%d %H:%M:%S") if lt else "",
            "detail": "%.1f credits used, last %dh ago" % (
                g["credits"], int(hours) if hours < 1000 else 999),
            "src": "usage",
        })
    tasks.sort(key=lambda t: t.get("date") or "", reverse=True)
    # cap: top MAX_RUNNING running + last MAX_DONE done (most recently finished)
    run = [t for t in tasks if t["status"] == "running"][:MAX_RUNNING]
    done = [t for t in tasks if t["status"] == "done"][:MAX_DONE]
    return run + done


if __name__ == "__main__":
    import json
    files = find_exports()
    print("export files found:", len(files))
    for f in files:
        print("  -", f)
    tasks = aggregate()
    print("tasks:", len(tasks))
    for t in tasks:
        print("  [%s] %-30s %6.2f credits  %s" % (
            t["status"], t["name"][:30], t["credits"], t["date"]))
    print(json.dumps(tasks, ensure_ascii=False, indent=1)[:600])
